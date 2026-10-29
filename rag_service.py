import hashlib
import json
import os
import urllib.error
import urllib.request
from io import BytesIO
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation
import xlrd

from database import connect


ALLOWED_DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".xls", ".pptx", ".txt", ".md", ".csv", ".json", ".html", ".htm", ".xml"}


def _ollama(path, payload):
    base_url = os.getenv("OLLAMA_BASE_URL", "http://127.0.0.1:11434").rstrip("/")
    request = urllib.request.Request(f"{base_url}{path}", data=json.dumps(payload).encode(), headers={"Content-Type": "application/json"})
    try:
        with urllib.request.urlopen(request, timeout=int(os.getenv("OLLAMA_TIMEOUT", "600"))) as response:
            return json.loads(response.read())
    except (urllib.error.URLError, TimeoutError) as exc:
        raise RuntimeError(f"No fue posible conectar con Ollama en {base_url}: {exc}") from exc


def _embed(texts):
    return _ollama("/api/embed", {"model": os.getenv("OLLAMA_EMBED_MODEL", "nomic-embed-text"), "input": texts})["embeddings"]


def _chunks(text, size=1200, overlap=200):
    clean = " ".join(text.split())
    return [clean[start:start + size] for start in range(0, len(clean), size - overlap) if clean[start:start + size].strip()]


def _vector(values):
    return "[" + ",".join(f"{float(value):.9g}" for value in values) + "]"


def ingest_pages(document_name, pages, raw=None):
    records = [(page, index, chunk) for page, text in pages for index, chunk in enumerate(_chunks(text))]
    if not records:
        raise ValueError("El documento no contiene texto indexable.")
    source = raw if raw is not None else "\n".join(text for _, text in pages).encode()
    digest = hashlib.sha256(source).hexdigest()
    embeddings = _embed([record[2] for record in records])
    with connect() as connection, connection.cursor() as cursor:
        cursor.execute("DELETE FROM sop_chunks WHERE document_name=%s", (document_name,))
        for (page, index, content), embedding in zip(records, embeddings):
            cursor.execute("""INSERT INTO sop_chunks
                (document_hash,document_name,page,chunk_index,content,embedding)
                VALUES (%s,%s,%s,%s,%s,%s::vector)""",
                (digest, document_name, page, index, content, _vector(embedding)))
        connection.commit()
    return len(records)


def ingest_text(document_name, text):
    return ingest_pages(document_name, [(1, text)], text.encode())


def extract_document(filename, raw):
    extension = Path(filename).suffix.lower()
    if extension not in ALLOWED_DOCUMENT_EXTENSIONS:
        raise ValueError(f"Tipo de archivo no compatible: {extension or 'sin extensión'}")
    stream = BytesIO(raw)
    if extension == ".pdf":
        return [(index + 1, page.extract_text() or "") for index, page in enumerate(PdfReader(stream).pages)]
    if extension == ".docx":
        return [(1, "\n".join(p.text for p in Document(stream).paragraphs))]
    if extension == ".xlsx":
        workbook = load_workbook(stream, read_only=True, data_only=True)
        return [(i + 1, f"Hoja: {sheet.title}\n" + "\n".join(" | ".join(str(v) for v in row if v is not None) for row in sheet.iter_rows(values_only=True))) for i, sheet in enumerate(workbook.worksheets)]
    if extension == ".xls":
        workbook = xlrd.open_workbook(file_contents=raw)
        return [(i + 1, f"Hoja: {sheet.name}\n" + "\n".join(" | ".join(str(sheet.cell_value(r, c)) for c in range(sheet.ncols)) for r in range(sheet.nrows))) for i, sheet in enumerate(workbook.sheets())]
    if extension == ".pptx":
        presentation = Presentation(stream)
        return [(i + 1, "\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text"))) for i, slide in enumerate(presentation.slides)]
    return [(1, raw.decode("utf-8", errors="replace"))]


def ingest_document(filename, raw):
    return ingest_pages(filename, extract_document(filename, raw), raw)


def list_documents():
    with connect() as connection, connection.cursor() as cursor:
        cursor.execute("""SELECT document_name,count(*) AS chunks,min(created_at) AS indexed_at
            FROM sop_chunks GROUP BY document_name ORDER BY document_name""")
        return cursor.fetchall()


def delete_document(document_name):
    with connect() as connection, connection.cursor() as cursor:
        cursor.execute("DELETE FROM sop_chunks WHERE document_name=%s RETURNING id", (document_name,))
        rows = cursor.fetchall()
        connection.commit()
        return len(rows)


def retrieve(question, limit=4):
    query = _vector(_embed([question])[0])
    with connect() as connection, connection.cursor() as cursor:
        cursor.execute("""SELECT document_name,page,content,embedding <=> %s::vector AS distance
            FROM sop_chunks ORDER BY embedding <=> %s::vector LIMIT %s""", (query, query, limit))
        return cursor.fetchall()


def answer_rag(question, mode="manuals"):
    chunks = retrieve(question)
    if not chunks:
        raise ValueError("No hay documentación indexada.")
    evidence, sources = [], []
    for chunk in chunks:
        label = f"{chunk['document_name']} · página {chunk['page']}"
        evidence.append(f"[{label}]\n{chunk['content']}")
        if label not in sources:
            sources.append(label)
    instruction = "Guía el diagnóstico paso a paso desde una parada segura y el aislamiento de energías. " if mode == "troubleshooting" else "Responde de forma precisa a partir del manual. "
    prompt = f"""{instruction}Responde en español usando exclusivamente la evidencia recuperada. No inventes datos.
Prioriza parada segura, control de energías, LOTO, HSE y las instrucciones del fabricante.
No autorices retirar guardas, anular protecciones ni intervenir equipos energizados.
Si falta evidencia, indícalo. Cita las fuentes entre corchetes. Máximo 220 palabras, texto plano.

EVIDENCIA:
{chr(10).join(evidence)}

PREGUNTA: {question}"""
    result = _ollama("/api/chat", {"model": os.getenv("OLLAMA_CHAT_MODEL", "llama3.1:8b"), "messages": [
        {"role": "system", "content": "Eres un copiloto documental de mantenimiento para extracción y manejo de carbón. No sustituyes al personal autorizado."},
        {"role": "user", "content": prompt}], "stream": False, "options": {"temperature": 0.1, "num_ctx": 4096, "num_predict": 320}})
    return result["message"]["content"], sources


def resolve_mode(question, requested="auto"):
    if requested != "auto":
        return requested
    normalized = question.casefold()
    predictive_terms = ("riesgo", "predictiv", "vibración", "vibracion", "temperatura", "corriente", "rpm", "presión", "presion", "telemetría", "telemetria", "lectura", "máquina crítica", "maquina critica")
    troubleshooting_terms = ("alarma", "falla", "fallo", "no arranca", "detenida", "atasco", "ruido", "diagnóstico", "diagnostico")
    if any(term in normalized for term in predictive_terms):
        return "predictive"
    if any(term in normalized for term in troubleshooting_terms):
        return "troubleshooting"
    return "manuals"


def answer_predictive(question):
    from database import latest_machine_readings

    rows = latest_machine_readings()
    if not rows:
        raise ValueError("No existen lecturas predictivas disponibles en PostgreSQL.")
    normalized = question.casefold()
    selected = [row for row in rows if row["id"].casefold() in normalized or row["name"].casefold() in normalized]
    selected = selected or sorted(rows, key=lambda row: float(row["risk"]), reverse=True)[:5]
    counts = {state: sum(row["status"] == state for row in rows) for state in ("healthy", "warning", "critical")}
    average = round(sum(float(row["risk"]) for row in rows) / len(rows), 1)
    context = [f"Resumen de {len(rows)} máquinas: {counts['healthy']} saludables, {counts['warning']} en alerta y {counts['critical']} críticas. Riesgo promedio {average}%."]
    sources = []
    for row in selected[:5]:
        recorded_at = row["recorded_at"].isoformat()
        context.append(
            f"{row['id']} {row['name']} ({row['area']}), lectura {recorded_at}: riesgo {float(row['risk']):.1f}%, estado {row['status']}; "
            f"vibración {float(row['vibration']):.2f} mm/s de máximo {float(row['max_vibration']):.2f}; temperatura {float(row['temperature']):.1f} °C de máximo {float(row['max_temperature']):.1f}; "
            f"corriente {float(row['current']):.1f} A frente a nominal {float(row['rated_current']):.1f}; RPM {float(row['rpm']):.0f} frente a nominal {float(row['nominal_rpm']):.0f}; "
            f"presión {float(row['pressure']):.2f} bar de máxima {float(row['max_pressure']):.2f}; uso {float(row['usage_hours']):.1f} h de máximo {float(row['max_usage_hours']):.1f}. Explicación calculada: {row['explanation']}"
        )
        sources.append(f"PostgreSQL · machine_readings · {row['id']} · {recorded_at}")
    manual_chunks = retrieve(question, limit=3)
    evidence = []
    for chunk in manual_chunks:
        label = f"{chunk['document_name']} · página {chunk['page']}"
        evidence.append(f"[{label}]\n{chunk['content']}")
        if label not in sources:
            sources.append(label)
    prompt = f"""Responde en español usando solo los datos actuales y los manuales suministrados. Distingue lectura, límite de referencia, interpretación y acción.
No inventes ni recalcules cifras. El riesgo es un índice demostrativo, no una probabilidad estadística validada. Prioriza parada segura y LOTO. Cita fuentes entre corchetes. Máximo 240 palabras, texto plano.

DATOS ACTUALES:
{chr(10).join(context)}

MANUALES RECUPERADOS:
{chr(10).join(evidence)}

PREGUNTA: {question}"""
    result = _ollama("/api/chat", {"model": os.getenv("OLLAMA_CHAT_MODEL", "llama3.1:8b"), "messages": [
        {"role": "system", "content": "Eres un copiloto de mantenimiento predictivo para minería de carbón. Sé fiel a la telemetría y documentación proporcionadas."},
        {"role": "user", "content": prompt}], "stream": False, "options": {"temperature": 0.0, "num_ctx": 4096, "num_predict": 380}})
    return result["message"]["content"], sources


def answer_chat(question, requested_mode="auto"):
    mode = resolve_mode(question, requested_mode)
    if mode == "predictive":
        answer, sources = answer_predictive(question)
    else:
        answer, sources = answer_rag(question, mode)
    return answer, sources, mode
